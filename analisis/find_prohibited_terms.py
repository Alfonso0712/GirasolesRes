import docx
import re

doc = docx.Document("docs/DOCUMENTO_TECNICO_GIRASOLESRES_IEEE_v2.docx")
terms = ["plin", "tarjeta", "pos", "transferencia", "multicanal"]

print("=== PÁRRAFOS ===")
for i, p in enumerate(doc.paragraphs):
    for term in terms:
        if re.search(r"\b" + term, p.text, re.IGNORECASE):
            print(f"P{i} [{term}]: {p.text[:140]}...")
            break

print("\n=== TABLAS ===")
for t_idx, tbl in enumerate(doc.tables):
    for r_idx, row in enumerate(tbl.rows):
        for c_idx, cell in enumerate(row.cells):
            for term in terms:
                if re.search(r"\b" + term, cell.text, re.IGNORECASE):
                    txt = cell.text.replace("\n", " ")
                    print(f"T{t_idx+1} R{r_idx}C{c_idx} [{term}]: {txt[:140]}...")
                    break
